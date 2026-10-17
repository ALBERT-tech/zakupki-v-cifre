# -*- coding: utf-8 -*-
"""Список действий для Михаила по GetCourse (17.10.2026). python-docx. Вывод: docs/ и Downloads."""
import os, shutil
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "Михаилу_GetCourse_2026-10-17.docx")

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri"); st.paragraph_format.space_after = Pt(4)
for name, size in (("Title", 19), ("Heading 1", 13.5)):
    s = doc.styles[name]; s.font.name = "Calibri"; s.font.size = Pt(size); s.font.bold = True
    s.font.color.rgb = RGBColor(0x14, 0x16, 0x1D); s.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
doc.styles["Heading 1"].paragraph_format.space_before = Pt(14)
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = Cm(1.8); sec.left_margin = sec.right_margin = Cm(2)
MUTED = RGBColor(0x6B, 0x70, 0x7C); GREEN = RGBColor(0x0A, 0x6C, 0x40)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); e = OxmlElement("w:shd")
    e.set(qn("w:val"), "clear"); e.set(qn("w:color"), "auto"); e.set(qn("w:fill"), fill); tcPr.append(e)


def P(text, color=None, size=None, bold=False):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    return p


def TABLE(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5); shade(c, "E6EFE9")
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ""; r = cells[j].paragraphs[0].add_run(str(v)); r.font.size = Pt(9.5)
            cells[j].paragraphs[0].paragraph_format.space_after = Pt(0)
            if j == 0: r.bold = True; r.font.color.rgb = MUTED
    for row in t.rows:
        for j, w in enumerate(widths): row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc.add_heading("GetCourse школы: что сделано и что нужно от Михаила", level=0)
P("17 октября 2026. Источник истины — сайт cifrazakup.ru и Google-таблица расписания. "
  "Правки сделаны из учётной записи Альберта (роль «администратор»); перед каждой правкой страницы сделана резервная копия.", color=MUTED)

doc.add_heading("1. Уже сделано (проверять не нужно, только знать)", level=1)
TABLE(["№", "Что", "Где"], [
    ["1", "Главная школы gc.cifrazakup.ru — теперь клон cifrazakup.ru: ближайшее событие, расписание из таблицы, форма заявки, чат, вход в кабинет. "
          "Страница состоит из ОДНОГО блока, который подтягивает содержимое с сайта, поэтому обновляется сама. В редакторе GetCourse /main не править.", "/main"],
    ["2", "Публичная оферта заменена актуальной редакцией (г. Уфа, cifrazakup@yandex.ru) вместо июньской (Златоуст, личная почта).", "/oferta"],
    ["3", "Политика обработки ПДн заменена актуальной; исправлен заголовок вкладки (был «Публичная оферта»).", "/privacypolicy"],
    ["4", "Недописанный шаблон «Политика конфиденциальности» с заглушками [**] и личной почтой снят с публикации — он нигде не использовался.", "/privacypolicy1"],
    ["5", "Вебинар про претензионную работу: «30 дней» → «90 дней» (5 мест), доступ в предложении 8590447 — 90 дней вместо 10, "
          "служебная надпись «ссылки заполнить в GetCourse» в подвале заменена ссылками, «Сбербанк-АСТ» → «ЭТП «Сбер А»».", "/gzissues, предложение 8590447"],
    ["6", "Курс Марковой: третье занятие в форме оплаты подписано 17.12.2026 вместо 17.10.2026.", "предложение 8710743"],
    ["7", "Курс Милонаец: исправлен заголовок страницы («Практических курс» → «Практический курс»); в трёх предложениях латинская C в слове «Споры» заменена русской.",
          "/omilonaets_6steps, предложения 8756482, 8756483, 8814342"],
    ["8", "Вебинар Кельберг: в подвале официальный логотип Сбер А вместо знака с текстом. Форма регистрации не тронута.", "/ocenka_g_m_prop"],
    ["9", "Новая страница курса «Искусственный интеллект в закупках» для лицензии: программа по утверждённому PDF, кнопка на PDF, форма заявки на поток. "
          "Оплаты и предложений нет. Как и главная, это один блок, который подтягивает содержимое с cifrazakup.ru/course-ai.html. В редакторе не править.",
          "/course_ai (папка «Курсы»)"],
    ["10", "Резервные копии (не опубликованы): «Резерв главной до 17.10.2026», «Резерв оферты…», «Резерв политики ПДн…». Через месяц можно удалить.",
          "main_backup_1710, oferta_backup_1710, privacypolicy_backup_1710"],
], [0.8, 12.2, 4.0])
P("Отменяется прежняя задача «накатить главную семью блоками» (файл main-getcourse.html): главная уже стоит одним блоком.", color=GREEN, bold=True)

doc.add_heading("2. Нужно сделать Михаилу — это доступно только владельцу аккаунта", level=1)
TABLE(["№", "Что сделать", "Где в GetCourse", "Зачем"], [
    ["1", "Проверить реквизиты школы: ИП Герасимова Р. Г., г. Уфа, почта cifrazakup@yandex.ru. Убрать Златоуст и личную почту, если остались.",
          "Профиль → Настройки аккаунта", "Реквизиты подставляются в письма, счета и чеки"],
    ["2", "Проверить, на кого принимаются оплаты (ЮKassa / GetCourse Pay): получатель денег должен быть ИП Герасимова — лицензия и оферта на ИП.",
          "Настройки аккаунта → Платёжные системы / GetCourse Pay", "Деньги за образовательные услуги — исполнителю по оферте"],
    ["3", "Оформить двусторонний договор с GetCourse на ИП Герасимова: заполнить, отправить на проверку, подписать у Регины, скан — в поддержку. Срок 1 год.",
          "Профиль → Настройки аккаунта → Договор", "Нужен для лицензии на образовательную деятельность"],
    ["4", "Проверить отправителя и подпись писем (имя, адрес отправителя, подвал): без личной почты, с реквизитами ИП.",
          "Настройки аккаунта → Почта / Рассылки", "Единые контакты школы"],
    ["5", "Убедиться, что галочки согласия в формах ведут на /oferta и /privacypolicy (сейчас так и есть) и что включена отдельная галочка на рассылки.",
          "Настройки аккаунта → Согласия", "152-ФЗ и закон о рекламе"],
    ["6", "Выставлять счёт за тариф GetCourse и закрывающие документы так, как решили (компенсация ИП по договору); суммы — во вкладку «Траты».",
          "Профиль → Настройки аккаунта → Оплатить; вкладка «Акты (УПД)»", "Решение встречи 06.10"],
], [0.8, 7.4, 4.6, 4.2])

doc.add_heading("3. Нужно решение команды (технически может любой администратор)", level=1)
TABLE(["№", "Вопрос", "Варианты"], [
    ["1", "Вебинар Кельберг прошёл 06.10, а страница всё ещё зовёт «Записаться».", "Сменить призыв на «Получить запись» (страница записи ocenka_g_m_prop_zapis_0610 есть) или оставить регистрацию как способ получить запись"],
    ["2", "Промо-вебинар про претензии: страница /gzissues_promo снята, промо-предложение 8682221 (1 000 ₽, доступ 10 дней) актуально.", "Архивировать предложение или вернуть промо, исправив доступ на 90 дней и текст"],
    ["3", "Старое предложение «ИИ в госзакупках» за 30 000 ₽ (8502279). В таблице курс стоит 75 000 ₽.", "Страница курса /course_ai уже есть и работает без оплаты. Предложение архивировать или, когда будет решение о продаже, переименовать в «Искусственный интеллект в закупках» и поменять цену"],
    ["4", "Урок 3 Милонаец: в предложении 12:00, в таблице 12:30.", "Уточнить фактическое время и выровнять"],
    ["5", "Ссылка в таблице расписания у вебинара про претензии ведёт на выключенную /gzissues_promo.", "Поставить /gzissues (5 000 ₽) или оставить «Черновик»"],
], [0.8, 8.4, 7.8])

doc.add_heading("4. Правило на будущее", level=1)
P("Главную /main, страницу курса /course_ai, оферту /oferta и политику /privacypolicy в редакторе GetCourse не менять: их содержание берётся с cifrazakup.ru. "
  "Если нужно изменить текст — пишите Альберту, он поправит сайт, и GetCourse обновится (главная и курс — сами, документы — по пересборке).")

doc.save(OUT)
dl = os.path.join(os.path.expanduser("~"), "Downloads", os.path.basename(OUT))
shutil.copyfile(OUT, dl)
print("saved", OUT, "and", dl)
